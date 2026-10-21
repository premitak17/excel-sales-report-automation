"""Tests for loading, validation, calculations, and workbook output."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from sales_report.report import create_sales_report
from sales_report.validation import clean_sales_data, load_sales_data

PROJECT_ROOT = Path(__file__).resolve().parents[1]


class SalesReportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.sample_path = PROJECT_ROOT / "data" / "sample_sales.csv"
        self.raw_data = load_sales_data(self.sample_path)

    def test_sample_data_flags_duplicates_and_invalid_rows(self) -> None:
        result = clean_sales_data(self.raw_data)

        self.assertEqual(len(result.valid), 5)
        self.assertEqual(len(result.duplicates), 1)
        self.assertEqual(len(result.rejected), 3)
        self.assertIn("missing or invalid date", result.rejected["issues"].iloc[0])
        self.assertIn("quantity must be greater than zero", result.rejected["issues"].iloc[1])
        self.assertIn("missing or invalid date", result.rejected["issues"].iloc[2])

    def test_missing_columns_fail_clearly(self) -> None:
        with self.assertRaisesRegex(ValueError, "Missing required column"):
            clean_sales_data(pd.DataFrame({"order_id": ["1"]}))

    def test_alternative_date_column_names_are_normalized(self) -> None:
        for date_column in ("date", "order_date", "Order Date", "Sales Date", "  sAlEs   dAtE  "):
            with self.subTest(date_column=date_column):
                data = self.raw_data.rename(columns={"date": date_column})
                result = clean_sales_data(data)

                self.assertEqual(len(result.valid), 5)
                self.assertIn("date", result.valid.columns)
                self.assertTrue(pd.api.types.is_datetime64_any_dtype(result.valid["date"]))

    def test_missing_date_column_is_not_accepted(self) -> None:
        data = self.raw_data.drop(columns="date")

        with self.assertRaisesRegex(ValueError, "Missing required column\\(s\\): date"):
            clean_sales_data(data)

    def test_report_has_expected_sales_totals_and_sheets(self) -> None:
        result = clean_sales_data(self.raw_data)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "report.xlsx"
            metrics = create_sales_report(result, output_path)

            self.assertTrue(output_path.is_file())
            self.assertEqual(metrics["valid_rows"], 5)
            self.assertEqual(metrics["duplicate_rows"], 1)
            self.assertEqual(metrics["rejected_rows"], 3)
            self.assertAlmostEqual(metrics["total_revenue"], 657.5)

            workbook = load_workbook(output_path, data_only=True)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Summary",
                    "Monthly Sales",
                    "Top Products",
                    "Clean Data",
                    "Duplicates",
                    "Rejected Rows",
                ],
            )
            summary_values = {
                row[0].value: row[1].value
                for row in workbook["Summary"].iter_rows(min_row=2)
            }
            self.assertEqual(summary_values["Total revenue"], 657.5)
            self.assertEqual(workbook["Clean Data"].freeze_panes, "A2")
            self.assertEqual(workbook["Clean Data"]["G2"].number_format, '#,##0.00')
            top_products = workbook["Top Products"]
            self.assertEqual(top_products["A2"].value, "Widget A")
            self.assertEqual(top_products["B2"].value, 17)
            monthly_sales = workbook["Monthly Sales"]
            self.assertEqual(monthly_sales["A2"].value, "2025-01")
            self.assertEqual(monthly_sales["C2"].value, 250)

    def test_excel_input_is_supported(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_path = Path(temporary_directory) / "sales.xlsx"
            self.raw_data.to_excel(input_path, index=False)
            loaded = load_sales_data(input_path)

        self.assertEqual(len(loaded), len(self.raw_data))
        self.assertEqual(list(loaded.columns), list(self.raw_data.columns))


if __name__ == "__main__":
    unittest.main()