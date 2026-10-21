"""Sales calculations and professionally formatted Excel report creation."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sales_report.validation import ValidationResult


def create_sales_report(
    result: ValidationResult, output_path: str | Path
) -> dict[str, float | int]:
    """Write a multi-sheet Excel report and return its headline metrics."""
    output_file = Path(output_path)
    if output_file.suffix.lower() != ".xlsx":
        raise ValueError("The output report must use the .xlsx file extension.")
    output_file.parent.mkdir(parents=True, exist_ok=True)

    clean = result.valid.copy()
    clean["revenue"] = clean["quantity"] * clean["unit_price"]
    clean["date"] = pd.to_datetime(clean["date"])

    total_revenue = float(clean["revenue"].sum())
    total_orders = int(clean["order_id"].nunique())
    total_units = float(clean["quantity"].sum())
    monthly = (
        clean.assign(month=clean["date"].dt.to_period("M").astype(str))
        .groupby("month", as_index=False)
        .agg(total_units=("quantity", "sum"), total_revenue=("revenue", "sum"))
        .sort_values("month")
    )
    top_products = (
        clean.groupby("product", as_index=False)
        .agg(total_units=("quantity", "sum"), total_revenue=("revenue", "sum"))
        .sort_values(
            ["total_units", "total_revenue", "product"],
            ascending=[False, False, True],
        )
        .reset_index(drop=True)
    )
    summary = pd.DataFrame(
        {
            "Metric": [
                "Total revenue",
                "Unique orders",
                "Units sold",
                "Valid sales rows",
                "Duplicate rows",
                "Rejected rows",
            ],
            "Value": [
                total_revenue,
                total_orders,
                total_units,
                len(result.valid),
                len(result.duplicates),
                len(result.rejected),
            ],
        }
    )

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        monthly.to_excel(writer, sheet_name="Monthly Sales", index=False)
        top_products.to_excel(writer, sheet_name="Top Products", index=False)
        clean.to_excel(writer, sheet_name="Clean Data", index=False)
        result.duplicates.to_excel(writer, sheet_name="Duplicates", index=False)
        result.rejected.to_excel(writer, sheet_name="Rejected Rows", index=False)
        _format_workbook(writer)

    return {
        "total_revenue": total_revenue,
        "unique_orders": total_orders,
        "units_sold": total_units,
        "valid_rows": len(result.valid),
        "duplicate_rows": len(result.duplicates),
        "rejected_rows": len(result.rejected),
    }


def _format_workbook(writer: pd.ExcelWriter) -> None:
    workbook = writer.book
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_length = max(
                (len(str(cell.value)) if cell.value is not None else 0)
                for cell in column_cells
            )
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 40)

        header_to_column = {
            cell.value: cell.column for cell in worksheet[1] if cell.value is not None
        }
        for header in ("total_revenue", "unit_price", "revenue"):
            column = header_to_column.get(header)
            if column is not None:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=column).number_format = '#,##0.00'

        date_column = header_to_column.get("date")
        if date_column is not None:
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row=row, column=date_column).number_format = "yyyy-mm-dd"

    summary = workbook["Summary"]
    for row in range(2, summary.max_row + 1):
        if summary.cell(row=row, column=1).value == "Total revenue":
            summary.cell(row=row, column=2).number_format = '#,##0.00'